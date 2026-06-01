#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
通用 PPT 生成引擎
用法: python3 make_ppt.py slides_config.json
输出: PPTX + DALLE_prompts.md (放在桌面)
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ============================================================
# 配色常量
# ============================================================
C = {
    'DARK_BG':  RGBColor(0x0D,0x1B,0x2A), 'PRI_DARK': RGBColor(0x1B,0x3A,0x5C),
    'ACC_BLUE': RGBColor(0x2E,0x86,0xC1), 'LIT_BLUE': RGBColor(0xD4,0xE6,0xF1),
    'TXT_DARK': RGBColor(0x2C,0x3E,0x50), 'TXT_MED':  RGBColor(0x5D,0x6D,0x7E),
    'WHITE':    RGBColor(0xFF,0xFF,0xFF), 'LIT_BG':   RGBColor(0xF8,0xF9,0xFA),
    'WARM_ACC': RGBColor(0xE6,0x7E,0x22), 'GREEN_OK': RGBColor(0x27,0xAE,0x60),
    'RED_WARN': RGBColor(0xE7,0x4C,0x3C), 'V_LIGHT':  RGBColor(0xEB,0xED,0xEF),
    'WARM_BG':  RGBColor(0xFE,0xF5,0xE7), 'RED_BG':   RGBColor(0xFD,0xED,0xEC),
    'PURPLE':   RGBColor(0x8E,0x44,0xAD),
}

# ============================================================
# PPT 基础组件
# ============================================================
class SlideBuilder:
    def __init__(self, prs):
        self.prs = prs
        self.BL = prs.slide_layouts[6]

    def ns(self):
        return self.prs.slides.add_slide(self.BL)

    def tb(self, sl, l, t, w, h, tx='', fs=18, bd=False, cl=None, al=PP_ALIGN.LEFT):
        bx = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = bx.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = tx; p.font.size = Pt(fs)
        p.font.bold = bd; p.font.color.rgb = self._c(cl, 'TXT_DARK')
        p.font.name = 'PingFang SC'; p.alignment = al
        return tf

    def _c(self, v, default='LIT_BLUE'):
        """将颜色名或RGBColor统一转为RGBColor"""
        if isinstance(v, str): return C.get(v, C[default])
        return v or C[default]

    def rc(self, sl, l, t, w, h, fc=None, bc=None):
        s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb = self._c(fc, 'LIT_BLUE')
        if bc: s.line.color.rgb = self._c(bc, 'V_LIGHT'); s.line.width = Pt(0.5)
        else: s.line.fill.background()
        return s

    def hl(self, sl, l, t, w, cl=None, th=2):
        s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Pt(th))
        s.fill.solid(); s.fill.fore_color.rgb = self._c(cl, 'ACC_BLUE'); s.line.fill.background()

    def pn(self, sl, n):
        self.tb(sl, 12.2, 7.05, 1.0, 0.35, str(n), fs=10, cl=C['TXT_MED'], al=PP_ALIGN.RIGHT)

    def hd(self, sl, tx=''):
        self.tb(sl, 0.6, 0.25, 4.0, 0.35, tx, fs=9, cl=C['TXT_MED'])

    def ft(self, sl, tx=''):
        self.hl(sl, 0.6, 6.95, 12.1, C['V_LIGHT'], 1)
        if tx: self.tb(sl, 0.6, 7.02, 12.0, 0.35, tx, fs=8, cl=C['TXT_MED'])

    def ti(self, sl, tx, sub=None, ys=0.5):
        self.hl(sl, 0.6, ys+0.02, 1.2, C['ACC_BLUE'], 3)
        self.tb(sl, 0.6, ys+0.2, 11.5, 0.65, tx, fs=30, bd=True, cl=C['PRI_DARK'])
        if sub: self.tb(sl, 0.6, ys+0.8, 11.5, 0.45, sub, fs=13, cl=C['TXT_MED'])

    def card(self, sl, l, t, w, h, title, lines, icon='', color=None):
        c = self._c(color, 'ACC_BLUE')
        self.rc(sl, l, t, w, h, C['LIT_BG']); self.rc(sl, l, t, w, 0.06, c)
        if icon: self.tb(sl, l+0.25, t+0.2, 0.5, 0.5, icon, fs=26, bd=True, cl=c)
        self.tb(sl, l+0.25, t+0.55, w-0.5, 0.45, title, fs=17, bd=True, cl=C['PRI_DARK'])
        y = t + 1.05
        for ln in lines:
            if ln: self.tb(sl, l+0.25, y, w-0.5, 0.32, ln, fs=12, cl=C['TXT_DARK']); y += 0.3

    def info_box(self, sl, l, t, w, h, lines, bg=None, tc=None, fs=14):
        self.rc(sl, l, t, w, h, self._c(bg, 'LIT_BLUE'))
        self.tb(sl, l+0.4, t+0.15, w-0.8, h-0.3, '\n'.join(lines), fs=fs, cl=self._c(tc, 'TXT_DARK'))

    def table(self, sl, l, t, w, h, headers, rows, col_widths):
        nr = len(rows) + 1; nc = len(headers)
        tbl = sl.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(h)).table
        if col_widths:
            for i, cw in enumerate(col_widths): tbl.columns[i].width = Inches(cw)
        for j, hdr in enumerate(headers):
            cell = tbl.cell(0, j); cell.text = hdr
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C['WHITE']
                p.font.name = 'PingFang SC'; p.alignment = PP_ALIGN.CENTER
            tcPr = cell._tc.get_or_add_tcPr(); sf = cell._tc.makeelement(qn('a:solidFill'), {})
            sc = cell._tc.makeelement(qn('a:srgbClr'), {'val': '1B3A5C'}); sf.append(sc); tcPr.append(sf)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = tbl.cell(i+1, j); cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(11); p.font.color.rgb = C['TXT_DARK']
                    p.font.name = 'PingFang SC'; p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
                if i % 2 == 0:
                    tcPr = cell._tc.get_or_add_tcPr(); sf = cell._tc.makeelement(qn('a:solidFill'), {})
                    sc = cell._tc.makeelement(qn('a:srgbClr'), {'val': 'F8F9FA'}); sf.append(sc); tcPr.append(sf)
        return tbl

    def dalle_tag(self, sl):
        self.info_box(sl, 0.8, 6.3, 11.5, 0.45,
                      ['🤖 推荐用 DALL-E / 即梦生成图替换 → 见 prompt 文件'],
                      C['WARM_BG'], C['WARM_ACC'], fs=11)

    def image_placeholder(self, sl, l, t, w, h, txt=''):
        self.rc(sl, l, t, w, h, C['WHITE'], C['V_LIGHT'])
        if txt: self.tb(sl, l, t + h/2 - 0.2, w, 0.4, txt, fs=10, cl=C['TXT_MED'], al=PP_ALIGN.CENTER)

# ============================================================
# 页面渲染器
# ============================================================
class PageRenderer:
    def __init__(self, sb, config):
        self.sb = sb; self.cfg = config
        self.dalle_prompts = []  # 收集需要 DALL-E 生成的页面

    def render_all(self):
        meta = self.cfg.get('meta', {})
        course = meta.get('course', '')
        slides = self.cfg.get('slides', [])
        for i, sdata in enumerate(slides):
            page_num = i + 1
            stype = sdata.get('type', 'content')
            method = getattr(self, f'render_{stype}', self.render_content)
            method(sdata, page_num, course)
        print(f'  Rendered {len(slides)} slides')
        return self.dalle_prompts

    def _prep(self, sdata, page_num, course):
        sl = self.sb.ns()
        self.sb.hd(sl, course)
        sub = sdata.get('subtitle', '')
        self.sb.ti(sl, sdata.get('title', ''), sub if sub else None)
        self.sb.pn(sl, page_num)
        return sl

    def _color(self, name):
        return C.get(name, C['ACC_BLUE'])

    # ---- cover ----
    def render_cover(self, sdata, page_num, course):
        sl = self.sb.ns()
        self.sb.rc(sl, 0, 0, 13.333, 7.5, C['DARK_BG'])
        self.sb.hl(sl, 1.5, 2.0, 1.5, C['ACC_BLUE'], 4)
        self.sb.tb(sl, 1.5, 2.2, 10.0, 1.0, sdata.get('title',''), fs=44, bd=True, cl=C['WHITE'])
        st = sdata.get('subtitle', '')
        if st: self.sb.tb(sl, 1.5, 3.1, 10.0, 0.6, st, fs=20, cl=RGBColor(0xAe,0xBD,0xCA))
        self.sb.hl(sl, 1.5, 3.8, 3.0, RGBColor(0x34,0x49,0x5E), 1)
        self.sb.tb(sl, 1.5, 4.0, 10.0, 0.5, course, fs=14, cl=RGBColor(0x7F,0x8C,0x8D))
        icons = sdata.get('icons', [])
        for i, ic in enumerate(icons):
            self.sb.tb(sl, 2.0 + i*2.5, 5.4, 2.3, 0.5, ic, fs=15, cl=RGBColor(0x7F,0x8C,0x8D), al=PP_ALIGN.CENTER)
        ih = sdata.get('image_hint', '')
        if ih:
            self.sb.image_placeholder(sl, 9.5, 5.3, 2.8, 1.7, ih)

    # ---- content (default) ----
    def render_content(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        lines = sdata.get('lines', [])
        for lb in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *lb)
        for cd_item in sdata.get('cards', []):
            self.sb.card(sl, *cd_item)
        ih = sdata.get('image_hint', '')
        if ih: self.sb.image_placeholder(sl, 9.2, 1.8, 2.8, 2.0, ih)
        self.sb.ft(sl, sdata.get('footer', ''))

    # ---- cards (多卡片并列) ----
    def render_cards(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        cards = sdata.get('cards', [])
        n = len(cards)
        card_w = min(3.7, (11.5 - (n-1)*0.3) / n)
        for i, c in enumerate(cards):
            x = 0.8 + i * (card_w + 0.3)
            self.sb.card(sl, x, 1.6, card_w, c.get('h', 3.8), c['title'], c.get('lines', []),
                         c.get('icon', ''), self._color(c.get('color', 'ACC_BLUE')))
        for ibx in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *ibx)
        self.sb.ft(sl, sdata.get('footer', ''))

    # ---- table ----
    def render_table(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        tdata = sdata.get('table', {})
        headers = tdata.get('headers', [])
        rows = tdata.get('rows', [])
        widths = tdata.get('widths', [])
        if headers and rows:
            self.sb.table(sl, 0.8, 1.7, 11.5, min(len(rows)*0.9+0.5, 4.0), headers, rows, widths)
        for ibx in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *ibx)
        self.sb.ft(sl, sdata.get('footer', ''))

    # ---- layers (多层架构) ----
    def render_layers(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        layers = sdata.get('layers', [])
        for i, ly in enumerate(layers):
            y = 1.6 + i * 0.95
            self.sb.rc(sl, 0.8, y, 11.5, 0.85, C['LIT_BG'])
            self.sb.rc(sl, 0.8, y, 0.06, 0.85, self._color(ly.get('color', 'ACC_BLUE')))
            self.sb.tb(sl, 1.1, y+0.08, 3.5, 0.35, ly.get('label',''), fs=15, bd=True, cl=self._color(ly.get('color','ACC_BLUE')))
            self.sb.tb(sl, 4.8, y+0.1, 7.2, 0.35, ly.get('desc',''), fs=13, cl=C['TXT_DARK'])
        for ibx in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *ibx)
        self.sb.ft(sl, sdata.get('footer', ''))

    # ---- comparison (左右对比) ----
    def render_comparison(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        left = sdata.get('left', {})
        right = sdata.get('right', {})
        self.sb.rc(sl, 0.8, 1.6, 5.5, 0.55, C['RED_BG'])
        self.sb.tb(sl, 1.0, 1.63, 5.0, 0.45, left.get('header',''), fs=20, bd=True, cl=C['RED_WARN'])
        for i, item in enumerate(left.get('items', [])):
            y = 2.3 + i*1.15
            self.sb.rc(sl, 0.8, y, 5.5, 0.95, C['RED_BG'])
            self.sb.rc(sl, 0.8, y, 0.06, 0.95, C['RED_WARN'])
            self.sb.tb(sl, 1.1, y+0.05, 4.9, 0.4, item[0], fs=16, bd=True, cl=C['RED_WARN'])
            self.sb.tb(sl, 1.1, y+0.5, 4.9, 0.3, item[1] if len(item)>1 else '', fs=12, cl=C['TXT_DARK'])
        self.sb.rc(sl, 6.8, 1.6, 5.5, 0.55, C['LIT_BG'])
        self.sb.tb(sl, 7.0, 1.63, 5.0, 0.45, right.get('header',''), fs=20, bd=True, cl=C['GREEN_OK'])
        for i, item in enumerate(right.get('items', [])):
            y = 2.3 + i*1.15
            self.sb.rc(sl, 6.8, y, 5.5, 0.95, C['LIT_BG'])
            self.sb.rc(sl, 6.8, y, 0.06, 0.95, C['GREEN_OK'])
            self.sb.tb(sl, 7.1, y+0.15, 4.9, 0.5, item[0] if isinstance(item, list) else item, fs=15, bd=True, cl=C['GREEN_OK'])
        for ibx in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *ibx)
        self.sb.ft(sl, sdata.get('footer', ''))

    # ---- summary (编号问题总结) ----
    def render_summary(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        items = sdata.get('items', [])
        for i, item in enumerate(items):
            y = 1.8 + i*1.45
            circle = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(y+0.05), Inches(0.5), Inches(0.5))
            circle.fill.solid(); circle.fill.fore_color.rgb = self._color(item.get('color','ACC_BLUE'))
            circle.line.fill.background()
            self.sb.tb(sl, 0.8, y+0.08, 0.5, 0.5, str(i+1), fs=18, bd=True, cl=C['WHITE'], al=PP_ALIGN.CENTER)
            self.sb.tb(sl, 1.5, y, 4.0, 0.45, item.get('q',''), fs=22, bd=True, cl=C['PRI_DARK'])
            self.sb.tb(sl, 1.5, y+0.45, 8.0, 0.35, item.get('detail',''), fs=14, cl=C['TXT_MED'])
        self.sb.hl(sl, 0.8, 6.1, 11.5, C['V_LIGHT'], 1)
        nx = sdata.get('next', '')
        if nx: self.sb.tb(sl, 0.8, 6.2, 11.5, 0.4, nx, fs=14, bd=True, cl=C['ACC_BLUE'])
        for ibx in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *ibx)
        self.sb.ft(sl, sdata.get('footer', ''))

    # ---- dalle (标记为需要 AI 图像生成) ----
    def render_dalle(self, sdata, page_num, course):
        sl = self._prep(sdata, page_num, course)
        cards = sdata.get('cards', [])
        for c in cards:
            self.sb.card(sl, *c)
        for ibx in sdata.get('info_boxes', []):
            self.sb.info_box(sl, *ibx)
        ih = sdata.get('image_hint', '')
        if ih: self.sb.image_placeholder(sl, 9.2, 1.8, 2.8, 2.0, ih)
        self.sb.dalle_tag(sl)
        self.sb.ft(sl, sdata.get('footer', ''))
        # 收集 DALL-E prompt
        self.dalle_prompts.append({
            'page': page_num,
            'title': sdata.get('title', ''),
            'subtitle': sdata.get('subtitle', ''),
            'dalle_desc': sdata.get('dalle_desc', ''),
            'image_hint': ih,
        })

# ============================================================
# DALL-E Prompt 生成器
# ============================================================
def generate_dalle_md(prompts, meta):
    lines = [
        f"# DALL-E / 即梦 图像生成 Prompts — {meta.get('title', 'PPT')}",
        "",
        "> 逐条复制到 DALL-E / ChatGPT 图像生成 / 即梦 AI 中，生成 16:9 PNG 后插入 PPT 替换对应页。",
        "",
        "---",
        ""
    ]
    for p in prompts:
        lines.append(f"## 第 {p['page']} 页：{p['title']}")
        lines.append("")
        lines.append("```")
        lines.append(f"Create one complete 16:9 consulting concept-slide image.")
        lines.append("")
        lines.append(f"Slide role: evidence / concept")
        lines.append(f"Visual mode: Clear Report Exhibit")
        lines.append("")
        lines.append(f"Action title: {p['title']}")
        if p['subtitle']:
            lines.append(f"Subtitle: {p['subtitle']}")
        lines.append("")
        if p['dalle_desc']:
            lines.append(f"Main proof object: {p['dalle_desc']}")
        else:
            lines.append(f"Main proof object: {p.get('image_hint', 'Professional consulting diagram')}")
        lines.append("")
        lines.append("Style: High-authority consulting exhibit. White base, charcoal type, blue as anchor color.")
        lines.append("Light rules, generous margins, clear hierarchy. Chinese text legible.")
        lines.append("Avoid: Generic cards, stock photos, SaaS dashboard look, all-one-color, bullet points.")
        lines.append("")
        lines.append("Output: One fully generated full-slide PNG, 16:9.")
        lines.append("```")
        lines.append("")
        lines.append("---")
        lines.append("")
    return '\n'.join(lines)

# ============================================================
# 主入口
# ============================================================
def main():
    if len(sys.argv) < 2:
        print("用法: python3 make_ppt.py slides_config.json")
        sys.exit(1)

    config_path = sys.argv[1]
    with open(config_path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    meta = config.get('meta', {})
    output_name = meta.get('output_name', 'output')
    course = meta.get('course', '')

    # 创建 PPT
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)

    sb = SlideBuilder(prs)
    renderer = PageRenderer(sb, config)

    print(f'Building: {meta.get("title", "PPT")}')
    print(f'Slides: {len(config.get("slides", []))}')
    dalle_prompts = renderer.render_all()

    # 保存 PPTX
    pptx_path = os.path.expanduser(f'~/Desktop/{output_name}.pptx')
    prs.save(pptx_path)
    print(f'PPTX saved: {pptx_path}')

    # 生成 DALL-E prompt 文件
    dalle_count = len(dalle_prompts)
    if dalle_count > 0:
        md_path = os.path.expanduser(f'~/Desktop/{output_name}-DALLE-prompts.md')
        md_content = generate_dalle_md(dalle_prompts, meta)
        with open(md_path, 'w', encoding='utf-8') as f:
            f.write(md_content)
        print(f'DALL-E prompts ({dalle_count} pages): {md_path}')

    print(f'\nTotal: {len(config.get("slides",[]))} slides | {dalle_count} DALL-E pages | {len(config.get("slides",[])) - dalle_count} python-pptx pages')
    print('Done!')

if __name__ == '__main__':
    main()
