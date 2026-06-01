#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPT 参考文件读取器
用法: python3 read_refs.py <refs_dir> <brief.md>
输出: context_package.json（结构化上下文，喂给 Claude 生成 slides_config.json）
"""
import os, sys, json
from pptx import Presentation

def extract_pptx_text(filepath):
    """提取 PPTX 文件中所有文本"""
    try:
        prs = Presentation(filepath)
        slides_text = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t and len(t) > 2:
                            texts.append(t)
            if texts:
                slides_text.append({
                    'slide': i,
                    'title': texts[0][:80] if texts else '',
                    'content': '\n'.join(texts)
                })
        return {
            'file': os.path.basename(filepath),
            'slide_count': len(prs.slides),
            'slides': slides_text
        }
    except Exception as e:
        return {'file': os.path.basename(filepath), 'error': str(e)}

def parse_brief(filepath):
    """解析用户 brief 文件"""
    if not os.path.exists(filepath):
        return {}
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    result = {}
    current_key = None
    current_value = []

    for line in content.split('\n'):
        line = line.strip()
        if not line:
            continue
        if ':' in line and not line.startswith('-') and not line.startswith('#'):
            if current_key:
                result[current_key] = '\n'.join(current_value).strip()
            parts = line.split(':', 1)
            current_key = parts[0].strip().lower().replace(' ', '_')
            current_value = [parts[1].strip()] if len(parts) > 1 else ['']
        elif current_key:
            current_value.append(line)

    if current_key:
        result[current_key] = '\n'.join(current_value).strip()

    return result

def main():
    if len(sys.argv) < 2:
        print("用法: python3 read_refs.py <refs_directory> [brief.md]")
        print("输出: context_package.json")
        sys.exit(1)

    refs_dir = sys.argv[1]
    brief_path = sys.argv[2] if len(sys.argv) > 2 else None

    # 读取所有 PPTX 文件
    refs = []
    if os.path.isdir(refs_dir):
        for fname in sorted(os.listdir(refs_dir)):
            if fname.endswith(('.pptx', '.ppt')) and not fname.startswith('~'):
                fpath = os.path.join(refs_dir, fname)
                print(f'  Reading: {fname}')
                ref_data = extract_pptx_text(fpath)
                refs.append(ref_data)
                print(f'    {ref_data.get("slide_count", 0)} slides')

    # 构建上下文包
    package = {
        'references': refs,
        'reference_summary': []
    }

    # 为每个参考文件生成摘要
    for ref in refs:
        if 'error' in ref:
            continue
        summary = {
            'file': ref['file'],
            'slides': ref['slide_count'],
            'topics': []
        }
        for s in ref['slides'][:10]:  # 只看前10页的标题
            if s['title']:
                summary['topics'].append(s['title'][:100])
        package['reference_summary'].append(summary)

    # 读取 brief
    if brief_path:
        package['brief'] = parse_brief(brief_path)

    # 输出
    out_path = os.path.join(os.getcwd(), 'context_package.json')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(package, f, ensure_ascii=False, indent=2)

    print(f'\nContext package saved: {out_path}')
    print(f'References: {len(refs)} files')
    total_slides = sum(r.get('slide_count', 0) for r in refs)
    print(f'Total reference slides: {total_slides}')

if __name__ == '__main__':
    main()
